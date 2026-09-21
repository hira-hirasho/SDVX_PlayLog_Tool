from __future__ import annotations

from copy import copy
from pathlib import Path
from threading import Lock
from typing import Any

from loguru import logger
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from app.play_record import PlayRecord


class PlayLogExcelError(Exception):
    """プレイ履歴Excel操作に関する例外。"""


class PlayLogExcel:
    """SDVX PlayLogのプレイ履歴Excelを操作する。"""

    SHEET_NAME = "プレイ履歴"
    TABLE_NAME = "sdvx_play_log_table"

    HEADERS = [
        "プレイID",
        "プレイ日時",
        "曲名",
        "アーティスト名",
        "難易度",
        "LEVEL",
        "SCORE",
        "Δ SCORE",
        "EX SCORE",
        "Δ EX SCORE",
    ]

    _lock = Lock()

    def __init__(
        self,
        workbook_path: str | Path,
        *,
        retry_count: int = 3,
        retry_interval_seconds: float = 2.0,
    ) -> None:
        self.workbook_path = Path(
            workbook_path
        )

        self.retry_count = retry_count
        self.retry_interval_seconds = (
            retry_interval_seconds
        )

    def save_record(
        self,
        record: PlayRecord,
    ) -> bool:
        """
        プレイ履歴を既存Excelテーブルへ保存する。

        Returns:
            True:
                今回の呼び出しで新規追加した。

            False:
                同じプレイIDが既に存在しており、
                既存行をそのまま維持した。

        プレイIDをキーとして冪等に動作するため、
        同じrecordを複数回保存しても二重登録しない。
        """

        with self._lock:
            return self._save_record_locked(
                record
            )

    def _save_record_locked(
        self,
        record: PlayRecord,
    ) -> bool:
        last_error: Exception | None = None

        for attempt in range(
            1,
            self.retry_count + 1,
        ):
            try:
                return self._save_record_once(
                    record
                )

            except Exception as exc:
                last_error = exc

                if attempt >= self.retry_count:
                    break

                logger.warning(
                    "EXCEL: save failed "
                    f"(attempt="
                    f"{attempt}/{self.retry_count}): "
                    f"{exc}"
                )

                import time

                time.sleep(
                    self.retry_interval_seconds
                )

        raise PlayLogExcelError(
            "プレイ履歴の保存に失敗しました: "
            f"{self.workbook_path}"
        ) from last_error

    def _save_record_once(
        self,
        record: PlayRecord,
    ) -> bool:
        """Excelテーブルへの冪等保存を1回実行する。"""

        self._validate_workbook()

        workbook = load_workbook(
            self.workbook_path
        )

        try:
            worksheet = self._get_worksheet(
                workbook
            )

            table = self._get_table(
                worksheet
            )

            headers = self._get_table_headers(
                worksheet,
                table,
            )

            self._validate_headers(
                headers
            )

            (
                start_row,
                end_row,
                start_col,
                end_col,
            ) = self._get_table_range(
                table
            )

            if len(headers) != len(
                self.HEADERS
            ):
                raise PlayLogExcelError(
                    "Excelテーブルの列数が仕様と一致しません。"
                )

            # --------------------------------------------------------
            # 既存プレイIDを検索
            # --------------------------------------------------------

            existing_row = self._find_play_id_row(
                worksheet=worksheet,
                start_row=start_row,
                end_row=end_row,
                play_id_column=start_col,
                play_id=record.play_id,
            )

            if existing_row is not None:
                logger.info(
                    "EXCEL: play record already exists; "
                    "skipping duplicate: play_id={}",
                    record.play_id,
                )

                return False

            # --------------------------------------------------------
            # 新規行の決定
            # --------------------------------------------------------

            target_row: int | None = None

            for row in range(
                start_row + 1,
                end_row + 1,
            ):
                if self._is_row_empty(
                    worksheet,
                    row,
                    start_col,
                    end_col,
                ):
                    target_row = row
                    break

            if target_row is None:
                target_row = end_row + 1

                table.ref = (
                    f"{get_column_letter(start_col)}"
                    f"{start_row}:"
                    f"{get_column_letter(end_col)}"
                    f"{target_row}"
                )

            values: list[Any] = [
                record.play_id,
                record.played_at,
                record.song_name,
                record.artist,
                record.difficulty,
                record.level,
                record.score,
                record.score_delta,
                record.ex_score,
                record.ex_score_delta,
            ]

            if len(values) != len(headers):
                raise PlayLogExcelError(
                    "書き込む値の数とExcelテーブルの列数が一致しません。"
                )

            # --------------------------------------------------------
            # 値を書き込む
            # --------------------------------------------------------

            for offset, value in enumerate(
                values
            ):
                worksheet.cell(
                    row=target_row,
                    column=start_col + offset,
                ).value = value

            # --------------------------------------------------------
            # 直前行の書式をコピー
            # --------------------------------------------------------

            if target_row > start_row + 1:
                source_row = target_row - 1
            else:
                source_row = start_row

            for column in range(
                start_col,
                end_col + 1,
            ):
                source_cell = worksheet.cell(
                    row=source_row,
                    column=column,
                )

                target_cell = worksheet.cell(
                    row=target_row,
                    column=column,
                )

                target_cell.font = copy(
                    source_cell.font
                )

                target_cell.fill = copy(
                    source_cell.fill
                )

                target_cell.border = copy(
                    source_cell.border
                )

                target_cell.alignment = copy(
                    source_cell.alignment
                )

                target_cell.number_format = (
                    source_cell.number_format
                )

                target_cell.protection = copy(
                    source_cell.protection
                )

            workbook.save(
                self.workbook_path
            )

            logger.info(
                "EXCEL: play record written: play_id={}",
                record.play_id,
            )

            return True

        finally:
            workbook.close()

    def exists(
        self,
        play_id: str,
    ) -> bool:
        """指定したプレイIDのプレイ記録が存在するか確認する。"""

        with self._lock:
            self._validate_workbook()

            workbook = load_workbook(
                self.workbook_path,
                read_only=True,
            )

            try:
                worksheet = self._get_worksheet(
                    workbook
                )

                table = self._get_table(
                    worksheet
                )

                headers = self._get_table_headers(
                    worksheet,
                    table,
                )

                self._validate_headers(
                    headers
                )

                (
                    start_row,
                    end_row,
                    start_col,
                    _,
                ) = self._get_table_range(
                    table
                )

                return (
                    self._find_play_id_row(
                        worksheet=worksheet,
                        start_row=start_row,
                        end_row=end_row,
                        play_id_column=start_col,
                        play_id=play_id,
                    )
                    is not None
                )

            finally:
                workbook.close()

    @staticmethod
    def _find_play_id_row(
        *,
        worksheet,
        start_row: int,
        end_row: int,
        play_id_column: int,
        play_id: str,
    ) -> int | None:
        """指定したプレイIDの行番号を検索する。"""

        for row in range(
            start_row + 1,
            end_row + 1,
        ):
            value = worksheet.cell(
                row=row,
                column=play_id_column,
            ).value

            if value is not None and str(value) == play_id:
                return row

        return None

    def _validate_workbook(self) -> None:
        if not self.workbook_path.exists():
            raise PlayLogExcelError(
                "Excelファイルが存在しません: "
                f"{self.workbook_path}"
            )

        if not self.workbook_path.is_file():
            raise PlayLogExcelError(
                "Excelパスがファイルではありません: "
                f"{self.workbook_path}"
            )

    def _get_worksheet(
        self,
        workbook,
    ):
        if self.SHEET_NAME not in workbook.sheetnames:
            raise PlayLogExcelError(
                "シートが存在しません: "
                f"{self.SHEET_NAME}"
            )

        return workbook[
            self.SHEET_NAME
        ]

    def _get_table(
        self,
        worksheet,
    ):
        if self.TABLE_NAME not in worksheet.tables:
            raise PlayLogExcelError(
                "Excelテーブルが存在しません: "
                f"{self.TABLE_NAME}"
            )

        return worksheet.tables[
            self.TABLE_NAME
        ]

    def _get_table_headers(
        self,
        worksheet,
        table,
    ) -> list[str]:
        (
            start_row,
            _,
            start_col,
            end_col,
        ) = self._get_table_range(
            table
        )

        return [
            worksheet.cell(
                row=start_row,
                column=column,
            ).value
            for column in range(
                start_col,
                end_col + 1,
            )
        ]

    def _validate_headers(
        self,
        headers: list[str],
    ) -> None:
        if headers != self.HEADERS:
            raise PlayLogExcelError(
                "Excelテーブルの列構成が仕様と一致しません.\n"
                f"expected={self.HEADERS}\n"
                f"actual={headers}"
            )

    @staticmethod
    def _get_table_range(
        table,
    ) -> tuple[int, int, int, int]:
        (
            min_col,
            min_row,
            max_col,
            max_row,
        ) = range_boundaries(
            table.ref
        )

        return (
            min_row,
            max_row,
            min_col,
            max_col,
        )

    @staticmethod
    def _is_row_empty(
        worksheet,
        row: int,
        start_col: int,
        end_col: int,
    ) -> bool:
        return all(
            worksheet.cell(
                row=row,
                column=column,
            ).value is None
            for column in range(
                start_col,
                end_col + 1,
            )
        )
